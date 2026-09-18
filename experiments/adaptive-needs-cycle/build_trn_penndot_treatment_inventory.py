"""Build the Pennsylvania existing-rumble-strip inventory pilot audit."""

from __future__ import annotations

import argparse
import hashlib
import json
from collections import Counter, defaultdict
from datetime import datetime
from decimal import Decimal
from pathlib import Path

import openpyxl


EXPECTED_HEADERS = [
    "District", "County", "County #", "State Route #", "Segment #",
    "Direction", "Segment Length (ft)", "# of Lanes", "Business Plan",
    "AADT", "Truck %", "Speed Limit (MPH)", "Begin Offset (ft)",
    "End Offset (ft)", "Location", "Type", "Pattern", "Test Date",
]
TREATMENT_TYPES = {"CL": "centerline", "SH": "shoulder", "EL": "edgeline", "TR": "transverse"}


def sha256(path: Path) -> str:
    return hashlib.sha256(path.read_bytes()).hexdigest()


def round_miles(feet: float) -> float:
    return round(feet / 5280, 6)


def union_feet(rows: list[dict], key_fields: tuple[str, ...]) -> tuple[int, float]:
    groups: dict[tuple, list[tuple[float, float]]] = defaultdict(list)
    for row in rows:
        begin, end = sorted((float(row["Begin Offset (ft)"]), float(row["End Offset (ft)"])))
        groups[tuple(row[field] for field in key_fields)].append((begin, end))
    total = 0.0
    for intervals in groups.values():
        intervals.sort()
        current_begin, current_end = intervals[0]
        for begin, end in intervals[1:]:
            if begin <= current_end:
                current_end = max(current_end, end)
            else:
                total += current_end - current_begin
                current_begin, current_end = begin, end
        total += current_end - current_begin
    return len(groups), total


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("--workbook", type=Path, required=True)
    parser.add_argument("--landing-page", type=Path, required=True)
    parser.add_argument("--hpms-state-totals", type=Path, required=True)
    parser.add_argument("--output", type=Path, required=True)
    args = parser.parse_args()

    workbook = openpyxl.load_workbook(args.workbook, read_only=True, data_only=True)
    if set(workbook.sheetnames) != {"tblRumbleStrips", "Definitions"}:
        raise ValueError(f"unexpected workbook sheets: {workbook.sheetnames}")
    sheet = workbook["tblRumbleStrips"]
    values = sheet.iter_rows(values_only=True)
    headers = list(next(values))
    if headers != EXPECTED_HEADERS:
        raise ValueError("unexpected PennDOT workbook headers")
    rows = [dict(zip(headers, values_row)) for values_row in values]
    if not rows:
        raise ValueError("PennDOT workbook contains no treatment rows")
    if any(row[header] is None for row in rows for header in headers):
        raise ValueError("PennDOT treatment table contains missing cells")

    landing = args.landing_page.read_text(encoding="utf-8-sig")
    if "more than 6,800 miles of centerline rumble strips as of September 2022" not in landing:
        raise ValueError("expected PennDOT public 6,800-mile statement not found")

    hpms_rows = json.loads(args.hpms_state_totals.read_text(encoding="utf-8-sig"))
    pennsylvania = [row for row in hpms_rows if int(row["stateid"]) == 42]
    if len(pennsylvania) != 1:
        raise ValueError("expected one Pennsylvania HPMS state-total row")
    hpms_pa = pennsylvania[0]

    line_feet = 0.0
    reverse_offset_rows = 0
    interval_exceeds_segment_rows = 0
    for row in rows:
        begin = float(row["Begin Offset (ft)"])
        end = float(row["End Offset (ft)"])
        segment = float(row["Segment Length (ft)"])
        if end < begin:
            reverse_offset_rows += 1
        interval = abs(end - begin)
        if interval > segment:
            interval_exceeds_segment_rows += 1
        line_feet += interval

    type_rows = Counter(str(row["Type"]) for row in rows)
    type_line_feet = Counter()
    for row in rows:
        type_line_feet[str(row["Type"])] += abs(
            float(row["End Offset (ft)"]) - float(row["Begin Offset (ft)"])
        )
    if set(type_rows) != set(TREATMENT_TYPES):
        raise ValueError(f"unexpected treatment types: {sorted(type_rows)}")

    date_years = Counter(
        row["Test Date"].year if isinstance(row["Test Date"], datetime) else int(str(row["Test Date"])[:4])
        for row in rows
    )
    two_way_two_lane = [
        row for row in rows
        if row["Direction"] == "B" and int(row["# of Lanes"]) == 2 and row["Type"] in {"CL", "SH", "EL"}
    ]
    footprint_keys = ("District", "County #", "State Route #", "Segment #", "Direction")
    footprint_segments, footprint_feet = union_feet(two_way_two_lane, footprint_keys)
    by_type = []
    for code, label in TREATMENT_TYPES.items():
        selected = [row for row in rows if row["Type"] == code]
        selected_segments, selected_union_feet = union_feet(selected, footprint_keys)
        by_type.append({
            "type_code": code,
            "type": label,
            "treatment_rows": len(selected),
            "summed_interval_line_miles": round_miles(type_line_feet[code]),
            "segment_direction_keys": selected_segments,
            "within_type_union_miles": round_miles(selected_union_feet),
        })

    output = {
        "record_id": "trn-penndot-existing-rumble-strip-inventory:2024:v1",
        "record_family": "trn_penndot_existing_rumble_strip_inventory",
        "version": "v1.draft",
        "status": "state_inventory_custodied_join_to_hpms_and_untreated_mileage_blocked",
        "as_of_date": "2026-07-28",
        "pilot_scope": "Pennsylvania road-owner inventory pilot",
        "source_custody": [
            {"path": str(path).replace("\\", "/"), "sha256": sha256(path)}
            for path in [args.workbook, args.landing_page, args.hpms_state_totals]
        ],
        "workbook_contract": {
            "sheets": workbook.sheetnames,
            "treatment_rows": len(rows),
            "districts": len({row["District"] for row in rows}),
            "counties": len({str(row["County"]).strip() for row in rows}),
            "required_headers": EXPECTED_HEADERS,
            "test_or_entry_year_rows": [
                {"year": year, "rows": date_years[year]} for year in sorted(date_years)
            ],
            "date_boundary": "The workbook defines Test Date as either original video-log collection or district data-entry date; it is not necessarily installation date.",
        },
        "inventory_measures": {
            "all_treatment_rows_summed_interval_line_miles": round_miles(line_feet),
            "by_treatment_type": by_type,
            "measurement_boundary": "Summed interval line-miles count separate centerline, shoulder, and edgeline records and are not roadway route-miles. Within-type unions remove interval overlap only inside each treatment type.",
        },
        "candidate_setting_triage": {
            "filter": "Direction=B (workbook definition: both/two-way traffic), # of Lanes=2, Type in (CL,SH,EL)",
            "treatment_rows": len(two_way_two_lane),
            "segment_direction_keys": footprint_segments,
            "union_interval_footprint_miles": round_miles(footprint_feet),
            "measurement_boundary": "The union removes overlap among treatment records sharing district, county, route, segment, and direction. It is an analytical treatment footprint, not official route-mile total, rural mileage, functional-system match, or untreated subtraction.",
        },
        "quality_audit": {
            "reverse_offset_rows_normalized_with_absolute_interval": reverse_offset_rows,
            "interval_exceeds_reported_segment_length_rows": interval_exceeds_segment_rows,
            "status": "anomalies_disclosed_not_silently_discarded",
        },
        "published_context_not_reconciled": {
            "value": 6800,
            "unit": "more_than_centerline_rumble_strip_miles",
            "as_of": "2022-09",
            "associated_claim": "PennDOT reports a 47 percent decline in head-on/opposite-direction sideswipe fatalities since 2002 alongside installation of more than 6,800 miles of centerline rumble strips.",
            "boundary": "The public statement is descriptive attribution, not a controlled causal estimate. Its mileage definition and coverage do not reconcile directly to workbook interval sums, so it is not substituted for the workbook-derived measures.",
        },
        "hpms_pennsylvania_comparison": {
            "state_id": 42,
            "rural_two_way_two_lane_filtered_rows": int(hpms_pa["rows"]),
            "rural_two_way_two_lane_filtered_section_miles": float(hpms_pa["miles"]),
            "subtraction_performed": False,
            "join_status": "blocked",
            "reason": "The PennDOT workbook lacks rural/urban and federal functional-system fields and a documented direct key to the HPMS spatial records. It may also differ in road ownership and coverage. Subtracting its footprint from HPMS would mix universes.",
        },
        "next_join_contract": {
            "owner_system": "PennDOT Roadway Management System and Location Reference System",
            "keys_available_in_workbook": ["district", "county_number", "state_route_number", "segment_number", "direction", "begin_offset_feet", "end_offset_feet"],
            "fields_required": [
                "urban_or_rural", "federal_functional_system", "facility_type",
                "through_lanes", "road_ownership", "pavement_condition_and_structure",
                "usable_shoulder_geometry", "bicycle_and_noise_constraints",
                "observed_target_crashes_by_severity_and_period",
            ],
            "required_result": "matched treated and untreated Pennsylvania rural two-way two-lane corridor inventory with coverage and nonmatch diagnostics",
        },
        "blocked_values": {
            "pennsylvania_comparable_treated_miles": None,
            "pennsylvania_untreated_design_eligible_miles": None,
            "national_untreated_design_eligible_miles": None,
            "target_crash_fatalities_and_injuries": None,
            "candidate_cost_billions": None,
            "candidate_outcome": None,
        },
        "claim_boundaries": {
            "official_state_inventory_raw_custodied": True,
            "state_inventory_structure_audited": True,
            "analytical_treatment_footprint_ready": True,
            "inventory_to_hpms_join_ready": False,
            "comparable_treated_mileage_ready": False,
            "untreated_mileage_ready": False,
            "published_decline_treated_as_causal": False,
            "national_extrapolation_allowed": False,
            "candidate_cost_ready": False,
            "allocation_ready": False,
            "funding_recommended": False,
            "public_release_authorized": False,
        },
    }
    args.output.write_text(json.dumps(output, indent=2) + "\n", encoding="utf-8")
    print(json.dumps({
        "output": str(args.output),
        "rows": len(rows),
        "two_way_two_lane_union_footprint_miles": round_miles(footprint_feet),
        "hpms_subtraction_performed": False,
    }))


if __name__ == "__main__":
    main()
